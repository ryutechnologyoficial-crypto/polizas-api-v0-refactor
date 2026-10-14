"""Lectura y escritura de pólizas en Excel (.xlsx).

El formato es el mismo para leer y para escribir, así que un archivo exportado se
puede volver a cargar:

  · hoja «Polizas»:    numero, asegurado, tipo, prima, fecha_inicio, fecha_fin
  · hoja «Siniestros»: numero_poliza, fecha, monto, descripcion, estado
  · hoja «Resumen»:    solo informativa (se ignora al leer)
"""
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from modelos import Poliza

COLUMNAS_POLIZA = ["numero", "asegurado", "tipo", "prima", "fecha_inicio", "fecha_fin"]
COLUMNAS_SINIESTRO = ["numero_poliza", "fecha", "monto", "descripcion", "estado"]
_CABECERA = Font(bold=True, color="FFFFFF")
_FONDO = PatternFill("solid", fgColor="1F4E8C")


def _hoja(libro: Workbook, titulo: str, columnas: list[str], filas: list[list], primera: bool = False):
    hoja = libro.active if primera else libro.create_sheet()
    hoja.title = titulo
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font, celda.fill = _CABECERA, _FONDO
    for fila in filas:
        hoja.append(fila)
    for i, col in enumerate(columnas, start=1):
        ancho = max([len(str(col))] + [len(str(f[i - 1])) for f in filas]) + 2
        hoja.column_dimensions[hoja.cell(row=1, column=i).column_letter].width = min(ancho, 45)
    hoja.freeze_panes = "A2"
    return hoja


def exportar(db: Session, ruta: str | Path) -> dict:
    """Escribe todas las pólizas y sus siniestros en `ruta` y devuelve cuántas filas escribió.

    Se escribe primero en un archivo temporal y luego se reemplaza el destino: si algo
    falla a mitad, el Excel anterior queda intacto.
    """
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    polizas = db.scalars(select(Poliza).options(selectinload(Poliza.siniestros)).order_by(Poliza.id)).all()

    filas_p = [[p.numero, p.asegurado, p.tipo, p.prima, p.fecha_inicio, p.fecha_fin] for p in polizas]
    filas_s = [[p.numero, s.fecha, s.monto, s.descripcion, s.estado] for p in polizas for s in p.siniestros]
    filas_r = [[p.numero, len(p.siniestros), round(sum(s.monto for s in p.siniestros), 2)] for p in polizas]

    libro = Workbook()
    hp = _hoja(libro, "Polizas", COLUMNAS_POLIZA, filas_p, primera=True)
    hs = _hoja(libro, "Siniestros", COLUMNAS_SINIESTRO, filas_s)
    hr = _hoja(libro, "Resumen", ["numero", "n_siniestros", "monto_total"], filas_r)
    for hoja, columnas_fecha, columnas_pesos in ((hp, "EF", "D"), (hs, "B", "C"), (hr, "", "C")):
        for letra in columnas_fecha:
            for celda in hoja[letra][1:]:
                celda.number_format = "yyyy-mm-dd"
        for letra in columnas_pesos:
            for celda in hoja[letra][1:]:
                celda.number_format = "#,##0.00"

    fd, temporal = tempfile.mkstemp(suffix=".xlsx", dir=ruta.parent)
    os.close(fd)
    try:
        libro.save(temporal)
        os.chmod(temporal, 0o644)  # mkstemp lo crea solo legible por su dueño
        os.replace(temporal, ruta)  # PermissionError si el archivo está abierto en Excel (Windows)
    finally:
        if os.path.exists(temporal):
            os.remove(temporal)
    return {"ruta": str(ruta.resolve()), "polizas": len(filas_p), "siniestros": len(filas_s)}


def _fecha(valor) -> str:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor).strip()


def _filas(hoja) -> list[dict]:
    filas = hoja.iter_rows(values_only=True)
    cabecera = [str(c).strip() if c is not None else "" for c in next(filas, [])]
    return [dict(zip(cabecera, fila)) for fila in filas if any(v is not None for v in fila)]


def leer(ruta: str | Path) -> list[dict]:
    """Devuelve las pólizas del Excel como diccionarios listos para `POST /polizas`."""
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        polizas = {}
        for f in _filas(libro["Polizas"]):
            numero = str(f["numero"]).strip()
            polizas[numero] = {
                "numero": numero,
                "asegurado": str(f["asegurado"]).strip(),
                "tipo": str(f["tipo"]).strip(),
                "prima": float(f["prima"]),
                "fecha_inicio": _fecha(f["fecha_inicio"]),
                "fecha_fin": _fecha(f["fecha_fin"]),
                "siniestros": [],
            }
        if "Siniestros" in libro.sheetnames:
            for f in _filas(libro["Siniestros"]):
                poliza = polizas.get(str(f["numero_poliza"]).strip())
                if poliza is not None:
                    poliza["siniestros"].append({
                        "fecha": _fecha(f["fecha"]),
                        "monto": float(f["monto"]),
                        "descripcion": str(f["descripcion"]).strip(),
                        "estado": str(f.get("estado") or "abierto").strip(),
                    })
        return list(polizas.values())
    finally:
        libro.close()
